#!/usr/bin/env python3
"""Editable PowerPoint of lecture 2, matching materials/Пара2/index.html."""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".pptx_deps"))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

FIG = ROOT / ".pptx_work"
OUT = ROOT / "output" / "pptx" / "Пара2_линейная_регрессия.pptx"

NAVY = RGBColor(0x10, 0x20, 0x44)
ACCENT = RGBColor(0x11, 0x68, 0xE8)
GOLD = RGBColor(0xE0, 0xB0, 0x00)
INK = RGBColor(0x33, 0x47, 0x5F)
PALE = RGBColor(0xF4, 0xF7, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD5, 0xE2, 0xF2)
YELLOW = RGBColor(0xFF, 0xF6, 0xD8)
GLOSS = RGBColor(0x17, 0x4B, 0x94)

W, H = 13.333, 7.5


def _font(run, size, bold, color):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def text(slide, l, t, w, h, content, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    lines = content.split("\n") if isinstance(content, str) else content
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        _font(run, size, bold, color)
    return box


def rect(slide, l, t, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PALE


def chrome(slide, eyebrow, heading, number):
    bg(slide)
    rect(slide, 0.55, 0.28, 0.46, 0.07, GOLD)
    text(slide, 0.55, 0.38, 11.5, 0.32, eyebrow.upper(), 13, True, ACCENT)
    size = 32 if len(heading) < 42 else 28 if len(heading) < 68 else 24
    height = 0.7 if size >= 30 and len(heading) < 42 else 1.15
    text(slide, 0.5, 0.68, 12.3, height, heading, size, True, NAVY)
    text(slide, 11.3, 7.12, 1.6, 0.28, f"{number} / 34", 12, False, INK, PP_ALIGN.RIGHT)
    return 0.68 + height + 0.12


def formula(slide, content, top, width=12.2, left=0.55, yellow=False, size=22):
    height = 0.62 + 0.32 * content.count("\n")
    rect(slide, left, top, width, height, YELLOW if yellow else WHITE)
    rect(slide, left, top, 0.08, height, GOLD if yellow else ACCENT)
    text(slide, left + 0.22, top + 0.08, width - 0.36, height - 0.1, content, size, True, NAVY)
    return top + height + 0.14


def rows(slide, items, top, label_w=3.3):
    y = top
    for label, body in items:
        text(slide, 0.55, y, label_w, 0.7, label, 18, True, ACCENT)
        text(slide, 0.55 + label_w + 0.15, y, 12.2 - label_w - 0.15, 0.72, body, 18, False, NAVY)
        rect(slide, 0.55, y + 0.68, 12.2, 0.015, LINE)
        y += 0.82
    return y


def mark(slide, content, top, width=12.2):
    height = 0.7 if len(content) < 90 else 0.95
    rect(slide, 0.55, top, width, height, YELLOW)
    rect(slide, 0.55, top, 0.08, height, GOLD)
    text(slide, 0.8, top + 0.1, width - 0.4, height - 0.12, content, 18, True, NAVY)
    return top + height + 0.12


def card(slide, l, t, w, h, heading, body):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LINE
    text(slide, l + 0.22, t + 0.16, w - 0.4, 0.4, heading, 18, True, ACCENT)
    text(slide, l + 0.22, t + 0.58, w - 0.4, h - 0.74, body, 16, False, NAVY)


def picture(slide, name, l, t, w, h):
    slide.shapes.add_picture(str(FIG / name), Inches(l), Inches(t), Inches(w), Inches(h))


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Пара 2. Линейная регрессия"
    blank = prs.slide_layouts[6]

    def new():
        return prs.slides.add_slide(blank)

    s = new()
    bg(s)
    rect(s, 0.55, 1.7, 0.46, 0.07, GOLD)
    text(s, 0.55, 1.9, 12, 0.35, "ДГТУ  ·  ВВЕДЕНИЕ В МАШИННОЕ ОБУЧЕНИЕ", 14, True, ACCENT)
    text(s, 0.5, 2.4, 12, 1.2, "Линейная регрессия", 54, True, NAVY)
    text(s, 0.55, 3.8, 11, 0.6, "Вторая лекция  ·  первое точное решение, семейство метрик и честный счёт", 22, False, INK)
    text(s, 11.3, 7.12, 1.6, 0.28, "1 / 34", 12, False, INK, PP_ALIGN.RIGHT)

    s = new(); y = chrome(s, "Вопрос пары", "Откуда берутся веса, если формулу можно решить точно?", 2)
    text(s, 0.55, y + 0.3, 12.2, 1.6, "И какое число после этого можно называть качеством, а не удачной арифметикой на тех же точках?", 28, False, NAVY)

    s = new(); y = chrome(s, "Откуда слово  ·  Фрэнсис Гальтон, 1885–1886", "Крайний рост родителей не переходит детям целиком", 3)
    rows(s, [
        ("Высокие родители", "сыновья тоже высокие, но в среднем ниже отцов"),
        ("Низкие родители", "сыновья тоже низкие, но в среднем выше отцов"),
        ("Средние родители", "сыновья остаются около среднего по популяции"),
    ], y + 0.15)
    text(s, 0.55, 6.15, 12.2, 0.8, "Гальтон назвал это regression towards mediocrity: возврат отклонения к середине. Отсюда имя метода.", 18, False, NAVY)

    s = new(); y = chrome(s, "Что именно легло на прямую", "Линия проходит через средние сыновей, и она положе 45°", 4)
    picture(s, "galton-line.svg.png", 0.4, y, 6.4, 4.4)
    formula(s, "E[y | x] = μ + β(x − μ)", y + 0.1, width=5.8, left=7.0, yellow=True, size=18)
    text(s, 7.15, y + 1.15, 5.6, 1.4, "При близких разбросах β — корреляция. Она меньше 1, если связь не идеальна.", 18, False, NAVY)
    text(s, 7.15, y + 2.6, 5.6, 1.6, "Мелкая точка — семья. Крупная — средний сын в этой группе. У Гальтона наклон около 2/3.", 18, False, INK)

    s = new(); y = chrome(s, "Что называют регрессией", "Ответ — число. Модель говорит, каким оно бывает при этих признаках", 5)
    rows(s, [
        ("Не класс", "регрессия не ставит метку «да / нет». Класс — другая задача, даже если в названии есть слово regression"),
        ("Условный ответ", "при тех же признаках число всё ещё случайно. Регрессия выбирает, какую сводку писать в прогноз"),
        ("Сегодняшняя сводка", "условное среднее, и оно линейно по весам:  E[y | x] = xᵀw"),
    ], y + 0.2)

    s = new(); y = chrome(s, "Какие бывают", "Четыре названия. Сегодня разбираем первые два", 6)
    table = s.shapes.add_table(5, 3, Inches(0.55), Inches(y + 0.1), Inches(12.2), Inches(3.6)).table
    table.columns[0].width = Inches(3.3)
    table.columns[1].width = Inches(4.6)
    table.columns[2].width = Inches(4.3)
    data = [
        ("Название", "Своими словами", "Один пример"),
        ("Простая линейная", "одно число по одному признаку", "20 + 2×площадь. 40 м² дают 100"),
        ("Множественная", "несколько признаков, одна цена", "20 + 2×площадь + 3×этаж = 115"),
        ("Полиномиальная", "среди признаков есть квадрат площади", "график гнётся. Считается так же, не сегодня"),
        ("Логистическая", "ответ не цена, а «купит / не купит»", "к этой лекции не относится"),
    ]
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    _font(run, 14, r == 0 or c == 0, WHITE if r == 0 else NAVY)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if r == 0 else WHITE
    text(s, 0.55, 6.35, 12.2, 0.5, "Первые две строки — эта пара. В обеих на выходе одно число.", 18, True, NAVY)

    s = new(); y = chrome(s, "Модель", "Прогноз — взвешенная сумма признаков и сдвиг", 7)
    y = formula(s, "ŷ  =  w₀ + w₁x₁ + … + w_d x_d  =  wᵀx", y + 0.15, yellow=True, size=26)
    rows(s, [
        ("wⱼ   вес", "насколько меняется прогноз при +1 к признаку"),
        ("w₀   сдвиг", "прогноз, когда все признаки равны нулю"),
        ("остаток", "y − ŷ,  что модель не объяснила"),
    ], y + 0.15)

    s = new(); y = chrome(s, "Матрица плана", "Столбец единиц делает сдвиг обычным весом", 8)
    y = formula(s, "ŷ = Xw      первый столбец X — единицы", y + 0.05, size=24)
    y = formula(s, "[ 1   40   5 ]  ·  [ 20,  2,  3 ]   =   1·20 + 40·2 + 5·3   =   115", y, yellow=True, size=22)
    text(s, 0.55, y + 0.1, 12.2, 1.1, "Одна строка — одна квартира: 40 м², 5-й этаж. Единица умножается на сдвиг 20. Без единицы нулевая площадь стоила бы ноль.", 20, False, INK)

    s = new(); y = chrome(s, "Категория как столбцы регрессии", "One-hot ставит одну единицу. One-cold ставит один ноль", 9)
    table = s.shapes.add_table(4, 3, Inches(0.55), Inches(y + 0.05), Inches(12.2), Inches(2.8)).table
    table.columns[0].width = Inches(2.6)
    table.columns[1].width = Inches(4.8)
    table.columns[2].width = Inches(4.8)
    for r, row in enumerate([
        ("Район", "one-hot", "one-cold"),
        ("Центр", "1   0   0", "0   1   1"),
        ("Север", "0   1   0", "1   0   1"),
        ("Юг", "0   0   1", "1   1   0"),
    ]):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    _font(run, 20, True, WHITE if r == 0 else NAVY)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if r == 0 else WHITE
    text(s, 0.55, 6.15, 12.2, 0.85, "One-cold — это единица минус one-hot. Коды 1, 2, 3 не являются ни тем ни другим: модель прочитает Юг как втрое больший Центр.", 18, False, INK)

    s = new(); y = chrome(s, "Те же три района", "Единица в one-hot стоит там, где в one-cold стоит ноль", 10)
    picture(s, "onehot.svg.png", 1.6, y, 10.1, 4.6)

    s = new(); y = chrome(s, "Что это делает с уравнением", "Оба кода вместе со сдвигом делают веса неединственными", 11)
    card(s, 0.5, y, 6.0, 2.7, "One-hot", "Сумма трёх столбцов повторяет столбец единиц. В каждой строке ровно одна единица.")
    card(s, 6.8, y, 6.0, 2.7, "One-cold", "В каждой строке две единицы, поэтому сумма столбцов равна двум столбцам единиц.")
    mark(s, "Сдвиг оставляют, один район выбрасывают. Его вес — разница с выброшенным районом.", y + 2.95)

    s = new(); y = chrome(s, "Линейная регрессия", "Нормальные уравнения", 12)
    formula(s, "XᵀX w  =  Xᵀy", y, width=7.3, yellow=True, size=26)
    rows(s, [
        ("Что минимизируем", "Q(w) = ‖y − Xw‖².  Делить на n место дна не двигает"),
        ("Градиент", "−2Xᵀ(y − Xw). Ноль производной и есть жёлтая формула"),
        ("Гессиан", "2XᵀX. Чаша открыта вверх. Один набор весов, если столбцы не копии"),
    ], y + 0.9, label_w=2.7)
    text(s, 0.55, 6.15, 7.4, 0.85, "Для прямой a + bx та же формула: суммы по x слева, суммы по y справа. На доске не решаем.", 14, False, INK)
    picture(s, "loss-steps.svg.png", 8.15, y + 0.05, 4.7, 4.9)

    s = new(); y = chrome(s, "Когда формула молчит", "Прогноз может быть единственным, а веса — нет", 13)
    rows(s, [
        ("Копия столбца", "one-hot или one-cold плюс сдвиг. Разные веса дают один и тот же прогноз"),
        ("Почти копия", "малый шум в ответах сильно качает веса. Прогноз при этом может почти не сдвинуться"),
        ("Как считают", "систему решают через SVD или QR. Обращение XᵀX раздувает неточность"),
    ], y + 0.25)

    s = new(); y = chrome(s, "Вероятностное чтение", "Квадрат — это гауссов шум. Модуль — лапласов", 14)
    card(s, 0.5, y, 6.0, 3.1, "Нормальные ошибки", "Промахи независимы и похожи на колокол. Самое подходящее место весов — минимум суммы квадратов.")
    card(s, 6.8, y, 6.0, 3.1, "Тяжёлые хвосты", "У распределения Лапласа самое подходящее место — минимум суммы модулей. Готовой системы уравнений нет.")
    text(s, 0.55, y + 3.35, 12.2, 0.8, "Квадрат тянет прямую к далёкой точке. Модуль смотрит скорее на то, где лежит большинство.", 20, True, NAVY)

    s = new(); y = chrome(s, "Когда весу можно верить", "Четыре условия", 15)
    rows(s, [
        ("Среднее на прямой", "по столбцам из таблицы"),
        ("Признак не из ошибки", "иначе вес съезжает"),
        ("Промахи похожи", "один размер, не пачками"),
        ("Столбцы разные", "иначе вес не произнести"),
    ], y + 0.1)
    mark(s, "Нормальный шум не нужен. Причину формула не даёт.", 6.15)

    s = new(); y = chrome(s, "Как читать один вес", "Плюс один к признаку сдвигает прогноз на его вес, если остальные столбцы не трогать", 16)
    card(s, 0.5, y + 0.1, 6.0, 3.3, "Что фраза значит", "Остальные столбцы зафиксированы. Сдвиг посчитан внутри той записи, которую выбрали.")
    card(s, 6.8, y + 0.1, 6.0, 3.3, "Чего она не значит", "Что признак вызывает ответ. Что вес устойчив, если столбцы ходят вместе. Что после деления на разброс вес всё ещё «рубли за метр».")

    s = new(); y = chrome(s, "Две работы", "Функция потерь ищет веса. Метрика судит решение", 17)
    card(s, 0.5, y, 6.0, 2.8, "Loss", "Её минимизируют на учебных строках. Квадрат удобен: есть наклон и готовая формула.")
    card(s, 6.8, y, 6.0, 2.8, "Metric", "Её считают на объектах, которые веса не учили. Она может быть модулем, долей или квантилем.")
    text(s, 0.55, y + 3.15, 12.2, 0.8, "Учить MSE и сообщать MAE — нормальная пара. Называть их одним числом — нет.", 22, True, NAVY)

    s = new(); y = chrome(s, "MSE", "Квадрат сильно штрафует крупный промах", 18)
    text(s, 0.55, y - 0.08, 8, 0.32, "средняя квадратичная ошибка, mean squared error", 16, True, GLOSS)
    formula(s, "MSE  =  (1/m) Σ (y − ŷ)²", y + 0.28, width=7.5, yellow=True, size=24)
    text(s, 0.55, y + 1.35, 7.5, 1.5, "Ошибки 1 и 3 входят как 1 и 9. Единицы — квадрат ответа. Лучшая константа — среднее учебных ответов.", 18, False, INK)
    text(s, 0.55, y + 2.9, 7.5, 1.2, "На трёх точках: у прямой около 0,06, у «всегда говори 1,67» около 0,22.", 18, True, NAVY)
    picture(s, "loss-mse.svg.png", 8.3, y + 0.15, 4.5, 4.2)

    s = new(); y = chrome(s, "RMSE", "Корень возвращает единицы ответа, но хвост остаётся главным", 19)
    text(s, 0.55, y - 0.05, 12, 0.32, "корень из средней квадратичной ошибки, root mean squared error", 16, True, GLOSS)
    formula(s, "RMSE  =  √MSE", y + 0.4, yellow=True, size=32)
    text(s, 0.55, y + 1.7, 12.2, 1.0, "Число снова можно назвать в рублях. Крупная ошибка всё равно вошла через квадрат, ещё до корня.", 22, False, INK)
    text(s, 0.55, y + 2.9, 12.2, 0.8, "RMSE — не средняя ошибка. Среднюю ошибку в тех же единицах даёт MAE.", 22, True, NAVY)

    s = new(); y = chrome(s, "MAE", "Модуль отвечает в тех же единицах, что и ответ", 20)
    text(s, 0.55, y - 0.08, 8, 0.32, "средняя абсолютная ошибка, mean absolute error", 16, True, GLOSS)
    formula(s, "MAE  =  (1/m) Σ |y − ŷ|", y + 0.28, width=7.5, yellow=True, size=24)
    text(s, 0.55, y + 1.35, 7.5, 1.4, "Ошибки 1, 2 и 3 дают MAE = 2. Лучшая константа — медиана, не среднее.", 18, False, INK)
    text(s, 0.55, y + 2.85, 7.5, 1.3, "На трёх точках: у прямой около 0,22, у среднего около 0,44. МНК этого выигрыша не гарантирует.", 18, True, NAVY)
    picture(s, "loss-mae.svg.png", 8.3, y + 0.15, 4.5, 4.2)

    s = new(); y = chrome(s, "Один и тот же хвост", "Квадрат тянет прямую к далёкой точке. Модуль почти не тянет", 21)
    picture(s, "outlier-lines.svg.png", 1.4, y, 10.5, 4.7)

    s = new(); y = chrome(s, "R²", "Ноль значит «не лучше среднего», не «ноль точности»", 22)
    text(s, 0.55, y - 0.08, 8, 0.32, "коэффициент детерминации, доля разброса", 16, True, GLOSS)
    formula(s, "R²  =  1 −  (сумма квадратов промахов) / (разброс ответов)", y + 0.28, width=7.4, yellow=True, size=16)
    text(s, 0.55, y + 1.4, 7.3, 1.6, "Левый столбик — разброс ответов вокруг своего среднего. Правый — то, что прямая не угадала. Правый ниже, поэтому доля больше нуля.", 18, False, INK)
    text(s, 0.55, y + 3.15, 7.3, 1.2, "На новых данных R² бывает отрицательным. В программе берут среднее того ответа, который передали.", 18, True, NAVY)
    picture(s, "r2-bars.svg.png", 8.15, y + 0.2, 4.7, 4.3)

    s = new(); y = chrome(s, "Adjusted R²", "Лишний столбец на учебных строках больше не бесплатен", 23)
    text(s, 0.55, y - 0.05, 12, 0.32, "скорректированный коэффициент детерминации", 16, True, GLOSS)
    formula(s, "R²adj  =  1 − (1 − R²) · (n − 1) / (n − p)", y + 0.45, yellow=True, size=28)
    text(s, 0.55, y + 1.8, 12.2, 1.1, "n — число строк, p — число столбцов вместе со сдвигом. Штраф уменьшает долю за каждый новый столбец.", 22, False, INK)
    text(s, 0.55, y + 3.1, 12.2, 1.0, "Это поправка на той же выборке, не замена проверки на других строках.", 22, True, NAVY)

    s = new(); y = chrome(s, "MAPE", "Маленький ответ весит больше большого", 24)
    text(s, 0.55, y - 0.05, 12, 0.32, "средняя абсолютная процентная ошибка, mean absolute percentage error", 16, True, GLOSS)
    formula(s, "MAPE  =  (1/m) Σ  |y − ŷ| / |y|", y + 0.45, yellow=True, size=28)
    text(s, 0.55, y + 1.8, 12.2, 1.2, "Промах делят на сам ответ. При ответе 0 делить не на что. На трёх точках: около 0,14 у прямой и около 0,33 у среднего.", 22, False, INK)
    text(s, 0.55, y + 3.2, 12.2, 1.0, "На слайде это доля, не проценты.", 22, True, NAVY)

    s = new(); y = chrome(s, "sMAPE", "Слово «симметричная» описывает формулу, не оптимум", 25)
    text(s, 0.55, y - 0.05, 12, 0.32, "симметричная средняя абсолютная процентная ошибка", 16, True, GLOSS)
    formula(s, "sMAPE  =  (1/m) Σ   2 |y − ŷ|  /  (|y| + |ŷ|)", y + 0.45, yellow=True, size=26)
    text(s, 0.55, y + 1.8, 12.2, 1.1, "Множитель 2 — договорённость. Встречается запись без него. В отчёте надо назвать, какая именно.", 22, False, INK)
    text(s, 0.55, y + 3.1, 12.2, 1.1, "При том же модуле ошибки больший прогноз раздувает знаменатель и уменьшает штраф.", 22, True, NAVY)

    s = new(); y = chrome(s, "RMSLE", "При том же модуле занижение весит больше завышения", 26)
    text(s, 0.55, y - 0.05, 12, 0.32, "корень из средней квадратичной логарифмической ошибки", 16, True, GLOSS)
    formula(s, "RMSLE  =  √ среднего  ( log(1+y) − log(1+ŷ) )²", y + 0.45, yellow=True, size=24)
    text(s, 0.55, y + 1.8, 12.2, 1.2, "Метрика сравнивает порядки, не рубли. Ответ 100: прогноз 50 дальше в логарифме, чем прогноз 150.", 22, False, INK)
    text(s, 0.55, y + 3.15, 12.2, 1.0, "Отрицательный ответ в логарифм не подставить. Нулевые цены живут за счёт единицы под логарифмом.", 20, True, NAVY)

    s = new(); y = chrome(s, "Pinball", "Квантиль τ штрафует недобор и перебор по-разному", 27)
    formula(s, "если промах ≥ 0:  штраф = τ · промах\nесли промах < 0:  штраф = (τ − 1) · промах", y, width=7.5, yellow=True, size=18)
    text(s, 0.55, y + 1.55, 7.5, 1.5, "При τ = 1/2 это половина модуля, лучшая константа — медиана. При τ > 1/2 недобор дороже.", 18, False, INK)
    text(s, 0.55, y + 3.15, 7.5, 1.1, "Оптимум по константе — квантиль уровня τ, не среднее.", 18, True, NAVY)
    picture(s, "loss-pinball.svg.png", 8.3, y, 4.5, 4.2)

    s = new(); y = chrome(s, "Хьюбер", "Квадрат рядом с нулём, линейный рост дальше порога", 28)
    formula(s, "до порога δ:  половина квадрата промаха\nдальше порога:  растёт как модуль", y, width=7.5, yellow=True, size=20)
    text(s, 0.55, y + 1.55, 7.5, 1.4, "В нуле функция гладкая, в отличие от модуля. Хвост не получает квадратичного голоса.", 18, False, INK)
    text(s, 0.55, y + 3.1, 7.5, 1.2, "Это способ учить, не число для отчёта. Порог выбирают на отложенных строках.", 18, True, NAVY)
    picture(s, "loss-huber.svg.png", 8.3, y, 4.5, 4.2)

    s = new(); y = chrome(s, "Где имеют право жить веса", "Уравнение решают на train. Судят на других строках", 29)
    rows(s, [
        ("Train", "столбцы, веса, среднее константы, среднее и разброс для масштаба"),
        ("Validation", "какой набор столбцов, какая метрика отчёта, какой порог у Хьюбера"),
        ("Test", "один отчёт уже выбранного решения. Правка после него отменяет отчёт"),
    ], y + 0.15)
    mark(s, "Среднее, масштаб и словарь районов считают только на train.", 6.05)

    s = new(); y = chrome(s, "Как резать", "Разрезов четыре. Честный тот, который повторяет будущий прогноз", 30)
    picture(s, "splits.svg.png", 0.7, y - 0.05, 12.0, 5.15)

    s = new(); y = chrome(s, "Где разрез ломает регрессию", "Уравнение видит строки, которых в бою не будет", 31)
    rows(s, [
        ("Без перемешивания", "таблица отсортирована по цене, и test оказывается другим режимом"),
        ("Время перемешали", "завтрашняя цена стоит в train и учит вес для вчера"),
        ("Группу разрезали", "один дом и в train, и в test. Прямая запоминает дом"),
        ("Словарь до разреза", "столбцы районов собраны по всей таблице, включая районы только из test"),
    ], y + 0.1)

    s = new(); y = chrome(s, "Что портит уравнение", "В матрицу попало то, чего не будет в момент прогноза", 32)
    rows(s, [
        ("Прокси ответа", "цена за метр, умноженная на площадь, восстанавливает цену. Веса будут красивыми и пустыми"),
        ("Статистика до разреза", "среднее пропуска, масштаб и словарь районов — только с train"),
        ("Одна сущность дважды", "сделки одного дома по обе стороны: модель запоминает дом, не зависимость"),
    ], y + 0.25)

    s = new(); y = chrome(s, "Проверка", "Какие утверждения верны?", 33)
    rows(s, [
        ("А", "Веса ищем шагами. На каждом шаге ошибка меньше"),
        ("Б", "Веса стоят там, где сумма квадратов промахов меньше всего"),
        ("В", "Район кодируют числами 1, 2, 3"),
    ], y + 0.2, label_w=1.0)
    text(s, 0.55, 6.2, 12.2, 0.6, "Ответ после паузы: верно А и Б. Коды 1, 2, 3 — не one-hot.", 20, True, ACCENT)

    s = new(); y = chrome(s, "Итог лекции", "Что прошли", 34)
    for i, line in enumerate([
        "Прогноз — сдвиг плюс веса признаков",
        "Веса ищем шагами, формула записывает минимум",
        "Район превращают в столбцы, один уровень выбрасывают",
        "Квадрат, модуль и доля разброса",
        "Учат на train, проверяют на других строках",
        "1 ноября смотрим задачу. Если ответ — класс, модель со следующей пары",
    ]):
        text(s, 0.55, y + i * 0.72, 12.2, 0.6, line, 22, False, NAVY)
        rect(s, 0.55, y + 0.55 + i * 0.72, 12.2, 0.015, LINE)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)
    print("slides", len(prs.slides))


if __name__ == "__main__":
    build()
